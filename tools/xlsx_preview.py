#!/usr/bin/env python3
"""Render an .xlsx as a spreadsheet-like HTML grid (for the store screenshot): xlsx_preview.py <in.xlsx> <out.html>.

Needs openpyxl (Odoo's venv has it). Durations ([h]:mm) and dates are shown as a spreadsheet app would."""
import html
import sys
from datetime import datetime, timedelta

import openpyxl
from openpyxl.utils import get_column_letter


def show(cell):
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, timedelta):
        minutes = round(v.total_seconds() / 60)
        return '%d:%02d' % (minutes // 60, minutes % 60)
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)) and 'h' in (cell.number_format or ''):
        minutes = round(v * 24 * 60)
        return '%d:%02d' % (minutes // 60, minutes % 60)
    return str(v)


def main(src, out):
    ws = openpyxl.load_workbook(src).active
    widths = [max(8, (ws.column_dimensions[get_column_letter(c)].width or 10)) * 7 for c in range(1, ws.max_column + 1)]
    rows = ['<tr><th></th>' + ''.join(f'<th style="width:{w}px">{get_column_letter(i + 1)}</th>' for i, w in enumerate(widths)) + '</tr>']
    for r in range(1, min(ws.max_row, 40) + 1):
        cells = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            style = []
            if cell.font and cell.font.b:
                style.append('font-weight:bold')
            if cell.font and cell.font.sz and cell.font.sz > 12:
                style.append('font-size:%dpx' % (cell.font.sz + 2))
            if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb not in (None, '00000000') and cell.fill.fill_type:
                style.append('background:#' + str(cell.fill.fgColor.rgb)[-6:])
            if isinstance(cell.value, (int, float, timedelta)):
                style.append('text-align:right')
            cells.append(f'<td style="{";".join(style)}">{html.escape(show(cell))}</td>')
        rows.append(f'<tr><th>{r}</th>' + ''.join(cells) + '</tr>')
    page = ('<!doctype html><meta charset="utf-8"><style>body{margin:0;font:13px "DejaVu Sans",Arial,sans-serif;background:#fff}'
            '.bar{background:#1d7044;color:#fff;padding:8px 12px;font-weight:bold}table{border-collapse:collapse}'
            'th{background:#f3f3f3;color:#555;font-weight:normal;border:1px solid #d4d4d4;padding:2px 6px}'
            'td{border:1px solid #e3e3e3;padding:3px 6px;white-space:nowrap;overflow:hidden;max-width:420px}</style>'
            f'<div class="bar">{html.escape(src.rsplit("/", 1)[-1])}</div><table>{"".join(rows)}</table>')
    with open(out, 'w', encoding='utf-8') as f:
        f.write(page)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
